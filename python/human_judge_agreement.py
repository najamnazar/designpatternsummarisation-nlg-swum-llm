#!/usr/bin/env python3
"""
human_judge_agreement.py - Agreement between the human annotator and the LLM fact-preservation judge

Reads the completed annotation workbook (design_pattern_preservation_annotation.xlsx, built by
misc/build_annotation_workbook.py), maps the summary letters A-D back to the tools with
evaluation-results/manual_annotation_key.csv, and compares every human label with the label the
LLM judge (python/judge_fact_preservation.py) gave the same fact for the same summary.

Steps:
    1. Deanonymise: write a copy of the workbook in which every summary letter also shows its tool.
       Formulas are replaced by the values Excel last calculated, so the copy is a fixed record; the
       script checks that no original cell changed.
    2. Join the 564 human labels (141 facts x 4 summaries) to fact_preservation_verdicts.csv on
       (project, design_pattern, file, method, fact_id), checking that the fact texts match.
    3. Agreement on labels: percentage agreement, Cohen's kappa and Gwet's AC1 for the three labels
       and for Preserved vs not Preserved, with 95% confidence intervals from a bootstrap over the
       30 files; confusion matrix; break-downs by tool, fact type, repository and pattern; a
       sensitivity check on the facts the annotator marked correct in the code.
    4. Agreement on scores: the preserved rate of each of the 120 summaries (Spearman, Pearson,
       ICC(A,1), mean difference), the tools' overall preserved rates and ranking, and the order of
       the four tools within each file (Kendall's tau-b, share of tool pairs ordered the same way).
    5. Annotator-only information (no judge counterpart): the code check of the facts, unsupported
       claims and the extra facts, by tool.

Inputs:
    design_pattern_preservation_annotation.xlsx      annotated workbook (values as last saved)
    evaluation-results/manual_annotation_key.csv     letter -> tool for F01-F30
    evaluation-results/fact_preservation_verdicts.csv judge labels (default corpora)

Outputs:
    design_pattern_preservation_annotation_deanonymised.xlsx
    evaluation-results/human_judge_agreement_labels.csv      one row per fact and summary
    evaluation-results/human_judge_agreement_metrics.csv     agreement statistics per scope
    evaluation-results/human_judge_agreement_confusion.csv   judge x human counts, overall and by tool
    evaluation-results/human_judge_agreement_summaries.csv   preserved rate of each summary, both raters
    evaluation-results/human_judge_agreement_tools.csv       label rates and rank of each tool, both raters
    evaluation-results/human_judge_agreement_per_file.csv    tool order agreement per file
    evaluation-results/human_judge_agreement_annotator_only.csv  code check, unsupported claims, extra facts
    evaluation-results/results.txt                           delimited block (replaced on re-run),
                                                             plain-text and LaTeX tables

Usage:
    .venv/bin/python python/human_judge_agreement.py
"""
from __future__ import annotations

import itertools
import re
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Sequence

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.formula import ArrayFormula
from scipy import stats
from sklearn.metrics import cohen_kappa_score

REPO_ROOT = Path(__file__).resolve().parent.parent
EVAL_DIR = REPO_ROOT / "evaluation-results"
ANNOTATION_XLSX = REPO_ROOT / "design_pattern_preservation_annotation.xlsx"
DEANON_XLSX = REPO_ROOT / "design_pattern_preservation_annotation_deanonymised.xlsx"
KEY_CSV = EVAL_DIR / "manual_annotation_key.csv"
VERDICTS_CSV = EVAL_DIR / "fact_preservation_verdicts.csv"
OUT_PREFIX = EVAL_DIR / "human_judge_agreement"
RESULTS_TXT = EVAL_DIR / "results.txt"

# Markers delimiting this script's block inside results.txt, so a re-run replaces it.
BLOCK_START = "<<<BEGIN HUMAN-JUDGE AGREEMENT>>>"
BLOCK_END = "<<<END HUMAN-JUDGE AGREEMENT>>>"

LETTERS = ["A", "B", "C", "D"]
LABELS = ["PRESERVED", "OMITTED", "INCORRECT"]
TOOLS = ["DPS_NLG", "DPS_LLM (GPT)", "DPS_LLM (Mistral)", "DPS_SWUM"]
FACT_TYPES = ["role", "kind", "inheritance", "constructor", "field", "method", "call"]
N_BOOT = 2000
SEED = 20261027

TOOL_FILL = PatternFill("solid", fgColor="E8EEF7")
TOOL_FONT = Font(name="Calibri", bold=True, color="1F3864")


# ----------------------------------------------------------------------------- workbook

def load_key() -> pd.DataFrame:
    return pd.read_csv(KEY_CSV, dtype=str)


def letter_to_tool(key: pd.DataFrame) -> Dict[str, Dict[str, str]]:
    return {r.file_id: {L: getattr(r, f"summary_{L}") for L in LETTERS} for r in key.itertuples()}


def find_row(ws, col: int, value: str, start: int = 1) -> int:
    for r in range(start, ws.max_row + 1):
        if ws.cell(r, col).value == value:
            return r
    raise ValueError(f"{ws.title}: no '{value}' in column {col}")


def read_annotations(key: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Facts x letters, the unsupported-claim row and the extra facts, from the file sheets."""
    wb = load_workbook(ANNOTATION_XLSX, data_only=True)
    tools = letter_to_tool(key)
    facts, claims = [], []
    for r in key.itertuples():
        ws = wb[r.file_id]
        header = find_row(ws, 2, "Fact type", start=10)
        row = header + 1
        while ws.cell(row, 1).value not in (None, "UC"):
            for i, L in enumerate(LETTERS):
                facts.append({
                    "file_id": r.file_id, "project": r.project, "design_pattern": r.design_pattern,
                    "file": r.file_name, "fact_id": int(float(ws.cell(row, 1).value)),
                    "fact_text_workbook": ws.cell(row, 3).value, "code_check": ws.cell(row, 4).value,
                    "code_check_note": ws.cell(row, 5).value, "letter": L, "method": tools[r.file_id][L],
                    "human_label": (ws.cell(row, 6 + i).value or "").strip().upper() or None,
                    "human_notes": ws.cell(row, 10).value,
                })
            row += 1
        if ws.cell(row, 1).value == "UC":
            for i, L in enumerate(LETTERS):
                claims.append({"file_id": r.file_id, "letter": L, "method": tools[r.file_id][L],
                               "unsupported_claim": ws.cell(row, 6 + i).value})
    ws = wb["Extra Facts"]
    extra = []
    for row in range(find_row(ws, 1, "#") + 1, ws.max_row + 1):
        fid = ws.cell(row, 2).value
        if not fid or not re.fullmatch(r"F\d\d", str(fid)):
            continue
        for i, L in enumerate(LETTERS):
            extra.append({"extra_id": int(float(ws.cell(row, 1).value)), "file_id": fid,
                          "fact_type": ws.cell(row, 4).value, "fact": ws.cell(row, 5).value, "letter": L,
                          "method": tools[fid][L], "label": ws.cell(row, 7 + i).value})
    return pd.DataFrame(facts), pd.DataFrame(claims), pd.DataFrame(extra)


def write_deanonymised(key: pd.DataFrame) -> int:
    """Copy of the workbook with the tool next to every letter; returns the number of cells checked."""
    tools = letter_to_tool(key)
    values = load_workbook(ANNOTATION_XLSX, data_only=True)
    wb = load_workbook(ANNOTATION_XLSX)
    for ws in wb.worksheets:  # freeze formulas to their last calculated values
        src = values[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, ArrayFormula) or (isinstance(c.value, str) and c.value.startswith("=")):
                    c.value = src[c.coordinate].value
    added: Dict[str, set] = {}

    def put(ws, coord: str, value, bold: bool = False):
        ws[coord] = value
        ws[coord].alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        if bold:
            ws[coord].font, ws[coord].fill = TOOL_FONT, TOOL_FILL
        added.setdefault(ws.title, set()).add(coord)

    for fid, mapping in tools.items():
        ws = wb[fid]
        first = find_row(ws, 1, "A", start=8)
        for i, L in enumerate(LETTERS):
            assert ws.cell(first + i, 1).value == L
            put(ws, f"K{first + i}", mapping[L], bold=True)
        header = find_row(ws, 2, "Fact type", start=10)
        for i, L in enumerate(LETTERS):
            cell = ws.cell(header, 6 + i)
            assert cell.value == f"Summary {L}", (fid, cell.value)
            cell.value = f"Summary {L}\n{mapping[L]}"
            added.setdefault(fid, set()).add(cell.coordinate)
        ws.column_dimensions["K"].width = 18
        ws.row_dimensions[header].height = 30

    ws = wb["All Annotations"]
    assert ws["J1"].value == "Summary" and ws["M1"].value is None
    put(ws, "M1", "Tool (deanonymised)", bold=True)
    for r in range(2, ws.max_row + 1):
        fid, L = ws.cell(r, 1).value, ws.cell(r, 10).value
        if fid in tools and L in LETTERS:
            put(ws, f"M{r}", tools[fid][L])
    ws.column_dimensions["M"].width = 20
    ws.auto_filter.ref = f"A1:M{ws.max_row}"

    ws = wb["Files"]
    header = find_row(ws, 1, "ID")
    for i, L in enumerate(LETTERS):
        put(ws, f"{'LMNO'[i]}{header}", f"Summary {L}", bold=True)
        ws.column_dimensions["LMNO"[i]].width = 18
    for r in range(header + 1, ws.max_row + 1):
        fid = ws.cell(r, 1).value
        if fid in tools:
            for i, L in enumerate(LETTERS):
                put(ws, f"{'LMNO'[i]}{r}", tools[fid][L])

    ws = wb["Extra Facts"]
    header = find_row(ws, 1, "#")
    put(ws, f"L{header}", "Tools of summaries A / B / C / D", bold=True)
    for r in range(header + 1, ws.max_row + 1):
        fid = ws.cell(r, 2).value
        if fid in tools:
            put(ws, f"L{r}", "\n".join(f"{L}: {tools[fid][L]}" for L in LETTERS))
    ws.column_dimensions["L"].width = 26

    wb.save(DEANON_XLSX)

    # Every cell of the original (as last calculated) must be unchanged apart from the added labels.
    new = load_workbook(DEANON_XLSX)
    checked = 0
    for ws in values.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.coordinate in added.get(ws.title, set()):
                    continue
                checked += 1
                if new[ws.title][c.coordinate].value != c.value:
                    raise AssertionError(f"{ws.title}!{c.coordinate} changed")
    return checked


# ----------------------------------------------------------------------------- statistics

def gwet_ac1(a: Sequence[str], b: Sequence[str], categories: Sequence[str]) -> float:
    a, b = np.asarray(a), np.asarray(b)
    pa = np.mean(a == b)
    pi = np.array([(np.mean(a == k) + np.mean(b == k)) / 2 for k in categories])
    pe = np.sum(pi * (1 - pi)) / (len(categories) - 1)
    return float((pa - pe) / (1 - pe)) if pe < 1 else float("nan")


def kappa(a, b, labels) -> float:
    if len(set(a) | set(b)) < 2:
        return float("nan")
    return float(cohen_kappa_score(a, b, labels=labels))


def binary(x: pd.Series) -> pd.Series:
    return np.where(x == "PRESERVED", "PRESERVED", "NOT_PRESERVED")


def label_stats(df: pd.DataFrame) -> Dict[str, float]:
    h, j = df.human.to_numpy(), df.judge.to_numpy()
    hb, jb = binary(df.human), binary(df.judge)
    return {
        "n": len(df), "agreement": float(np.mean(h == j)),
        "kappa": kappa(h, j, LABELS), "ac1": gwet_ac1(h, j, LABELS),
        "agreement_binary": float(np.mean(hb == jb)),
        "kappa_binary": kappa(hb, jb, ["PRESERVED", "NOT_PRESERVED"]),
        "ac1_binary": gwet_ac1(hb, jb, ["PRESERVED", "NOT_PRESERVED"]),
        "human_preserved": float(np.mean(h == "PRESERVED")), "judge_preserved": float(np.mean(j == "PRESERVED")),
        "human_incorrect": float(np.mean(h == "INCORRECT")), "judge_incorrect": float(np.mean(j == "INCORRECT")),
    }


def bootstrap_ci(df: pd.DataFrame, fn, rng: np.random.Generator) -> tuple[float, float]:
    """95% percentile interval, resampling whole files (labels within a file are not independent)."""
    groups = [g for _, g in df.groupby("file_id")]
    vals = []
    for _ in range(N_BOOT):
        sample = pd.concat([groups[i] for i in rng.integers(0, len(groups), len(groups))])
        v = fn(sample)
        if not np.isnan(v):
            vals.append(v)
    return (float(np.percentile(vals, 2.5)), float(np.percentile(vals, 97.5))) if vals else (np.nan, np.nan)


def icc_a1(x: np.ndarray, y: np.ndarray) -> float:
    """ICC(A,1): two-way, absolute agreement, single rater (McGraw and Wong 1996)."""
    data = np.column_stack([x, y])
    n, k = data.shape
    grand = data.mean()
    ms_r = k * np.sum((data.mean(axis=1) - grand) ** 2) / (n - 1)
    ms_c = n * np.sum((data.mean(axis=0) - grand) ** 2) / (k - 1)
    ss_e = np.sum((data - data.mean(axis=1, keepdims=True) - data.mean(axis=0, keepdims=True) + grand) ** 2)
    ms_e = ss_e / ((n - 1) * (k - 1))
    return float((ms_r - ms_e) / (ms_r + (k - 1) * ms_e + k * (ms_c - ms_e) / n))


def pair_order(a: Sequence[float], b: Sequence[float]) -> tuple[int, int, int]:
    """Tool pairs ordered the same way / the opposite way / tied by either rater."""
    same = opp = tie = 0
    for i, k in itertools.combinations(range(len(a)), 2):
        da, db = np.sign(a[i] - a[k]), np.sign(b[i] - b[k])
        if da == 0 or db == 0:
            tie += 1
        elif da == db:
            same += 1
        else:
            opp += 1
    return same, opp, tie


# ----------------------------------------------------------------------------- results.txt

def latex_table(metrics: pd.DataFrame, tools: pd.DataFrame) -> List[str]:
    def esc(s: str) -> str:
        return s.replace("_", r"\_")

    def ci(r) -> str:
        return f"[{r.kappa_ci_low:.2f}, {r.kappa_ci_high:.2f}]"

    lines = [
        r"\begin{table}[htbp]",
        r"  \centering",
        r"  \caption{Agreement between the human annotator and the LLM judge on fact preservation "
        r"(30 files, 141 facts, 4 summaries each).}",
        r"  \label{tab:human-judge-agreement}",
        r"  \begin{tabular}{lrrrrrrr}",
        r"    \toprule",
        r"    Scope & Labels & Agreement & $\kappa$ & 95\% CI & AC1 & Preserved (human) & Preserved (judge) \\",
        r"    \midrule",
    ]
    rate = tools.set_index("method")
    for _, r in metrics[metrics.scope.isin(["all", "sensitivity", "tool"])].iterrows():
        name = {"all labels": "All labels", "facts marked Correct in the code": "Facts correct in the code"}.get(
            r.value, r.value)
        if r.scope == "tool":
            hp, jp = rate.loc[r.value, "human_preserved_rate"], rate.loc[r.value, "judge_preserved_rate"]
        else:
            hp, jp = r.human_preserved, r.judge_preserved
        lines.append(f"    {esc(name)} & {r.n} & {r.agreement:.1%} & {r.kappa:.2f} & {ci(r)} & {r.ac1:.2f} & "
                     f"{hp:.1%} & {jp:.1%} \\\\".replace("%", r"\%"))
        if r.scope == "sensitivity":
            lines.append(r"    \midrule")
    lines += [r"    \bottomrule", r"  \end{tabular}", r"\end{table}"]
    return lines


def write_block(results_txt: Path, section: str) -> None:
    """Write the section into results.txt, replacing any previous block with the same markers."""
    existing = results_txt.read_text(encoding="utf-8") if results_txt.exists() else ""
    start, end = existing.find(BLOCK_START), existing.find(BLOCK_END)
    if start != -1 and end != -1:
        updated = existing[:start] + section + existing[end + len(BLOCK_END) + 1:]
    else:
        # blocks are separated by two blank lines
        updated = existing.rstrip("\n") + "\n\n\n" + section if existing else section
    results_txt.write_text(updated, encoding="utf-8")


# ----------------------------------------------------------------------------- analysis

def main() -> None:
    key = load_key()
    rng = np.random.default_rng(SEED)
    out: List[str] = []

    def say(line: str = "") -> None:
        out.append(line)
        print(line)

    checked = write_deanonymised(key)
    facts, claims, extra = read_annotations(key)
    assert facts.human_label.isin(LABELS).all(), facts[~facts.human_label.isin(LABELS)]

    verdicts = pd.read_csv(VERDICTS_CSV)
    merged = facts.merge(verdicts, on=["project", "design_pattern", "file", "method", "fact_id"],
                         how="left", validate="one_to_one")
    missing = merged.label.isna().sum()
    text_mismatch = (merged.fact_text_workbook.str.strip() != merged.fact_text.str.strip()).sum()
    assert missing == 0 and text_mismatch == 0, (missing, text_mismatch)
    df = merged.rename(columns={"human_label": "human", "label": "judge"})
    df["agree"] = df.human == df.judge

    cols = ["file_id", "project", "design_pattern", "file", "fact_id", "fact_type", "category", "fact_text",
            "code_check", "code_check_note", "letter", "method", "summary_words", "human", "judge", "agree",
            "evidence", "evidence_match", "justification", "human_notes"]
    df[cols].to_csv(f"{OUT_PREFIX}_labels.csv", index=False)

    say("=" * 78)
    say("HUMAN-JUDGE AGREEMENT (fact preservation: human annotation of 30 files vs the LLM judge)")
    say("=" * 78)
    say(f"Generated: {datetime.now():%Y-%m-%d %H:%M:%S}")
    say(f"Annotation: {ANNOTATION_XLSX.name} (letters A-D mapped with {KEY_CSV.name})")
    run = pd.read_csv(EVAL_DIR / "fact_preservation_results.csv")
    judge_model = run.loc[(run.section == "run") & (run.metric == "judge_model"), "value"].iloc[0]
    say(f"Judge: {judge_model} (default corpora, "
        "fact_preservation_verdicts.csv)")
    say(f"Files: {df.file_id.nunique()}   facts: {df.groupby(['file_id', 'fact_id']).ngroups}   "
        f"labels compared: {len(df)} (facts x 4 summaries)")
    say(f"Join: {len(df)} of {len(df)} labels matched; fact texts identical in workbook and judge output.")
    say(f"Deanonymised copy: {DEANON_XLSX.name} ({checked} original cells checked, 0 changed)")
    say()

    # --- label-level agreement
    rows = []

    def add(scope: str, value: str, sub: pd.DataFrame, ci: bool = True):
        s = {"scope": scope, "value": value, **label_stats(sub)}
        if ci and sub.file_id.nunique() >= 5:
            s["kappa_ci_low"], s["kappa_ci_high"] = bootstrap_ci(sub, lambda d: kappa(d.human, d.judge, LABELS), rng)
            s["kappa_binary_ci_low"], s["kappa_binary_ci_high"] = bootstrap_ci(
                sub, lambda d: kappa(binary(d.human), binary(d.judge), ["PRESERVED", "NOT_PRESERVED"]), rng)
        rows.append(s)

    add("all", "all labels", df)
    add("sensitivity", "facts marked Correct in the code", df[df.code_check == "Correct"])
    for t in TOOLS:
        add("tool", t, df[df.method == t])
    for t in FACT_TYPES:
        if (df.fact_type == t).any():
            add("fact_type", t, df[df.fact_type == t])
    for c in ["pattern role", "identity", "relationship", "responsibility"]:
        add("category", c, df[df.category == c])
    for p in ["AbdurRKhalid", "JamesZBL", "SpringFramework"]:
        add("repository", p, df[df.project == p])
    for p in sorted(df.design_pattern.unique()):
        add("pattern", p, df[df.design_pattern == p], ci=False)
    metrics = pd.DataFrame(rows)
    metrics.to_csv(f"{OUT_PREFIX}_metrics.csv", index=False, float_format="%.4f")

    def fmt_ci(r, k):
        lo, hi = r.get(f"{k}_ci_low"), r.get(f"{k}_ci_high")
        return f"[{lo:.2f}, {hi:.2f}]" if pd.notna(lo) else "".ljust(12)

    say("1. Label agreement (Preserved / Omitted / Incorrect)")
    say("-" * 78)
    say("kappa = Cohen's kappa; AC1 = Gwet's AC1 (less sensitive to skewed label use);")
    say("CI = 95% bootstrap interval over files. 'Binary' = Preserved vs Omitted+Incorrect.")
    say()
    say(f"{'Scope':<38}{'n':>5}{'Agree':>7}{'kappa':>7} {'CI':<13}{'AC1':>6}"
        f"{'Agr(b)':>8}{'k(b)':>6} {'CI(b)':<13}")
    for _, r in metrics.iterrows():
        name = r.value if r.scope in ("all", "sensitivity") else f"{r.scope}: {r.value}"
        say(f"{name[:37]:<38}{r.n:>5}{r.agreement:>7.0%}{r.kappa:>7.2f} {fmt_ci(r, 'kappa'):<13}{r.ac1:>6.2f}"
            f"{r.agreement_binary:>8.0%}{r.kappa_binary:>6.2f} {fmt_ci(r, 'kappa_binary'):<13}")
    say()

    # confusion matrices
    conf_rows = []
    say("Confusion matrix, all labels (rows = judge, columns = human)")
    for scope, sub in [("all", df)] + [(t, df[df.method == t]) for t in TOOLS]:
        cm = pd.crosstab(pd.Categorical(sub.judge, LABELS), pd.Categorical(sub.human, LABELS), dropna=False)
        for j in LABELS:
            for h in LABELS:
                conf_rows.append({"scope": scope, "judge": j, "human": h, "count": int(cm.loc[j, h])})
        if scope == "all":
            say(f"{'judge \\ human':<16}" + "".join(f"{h:>11}" for h in LABELS) + f"{'total':>8}")
            for j in LABELS:
                say(f"{j:<16}" + "".join(f"{cm.loc[j, h]:>11}" for h in LABELS) + f"{cm.loc[j].sum():>8}")
            say(f"{'total':<16}" + "".join(f"{cm[h].sum():>11}" for h in LABELS) + f"{cm.values.sum():>8}")
    pd.DataFrame(conf_rows).to_csv(f"{OUT_PREFIX}_confusion.csv", index=False)
    say()
    say("Per label, taking the human label as the reference:")
    for lab in LABELS:
        tp = ((df.judge == lab) & (df.human == lab)).sum()
        prec = tp / max((df.judge == lab).sum(), 1)
        rec = tp / max((df.human == lab).sum(), 1)
        say(f"  {lab:<10} judge precision {prec:5.0%} ({tp}/{(df.judge == lab).sum()}),"
            f" recall {rec:5.0%} ({tp}/{(df.human == lab).sum()})")
    hb, jb = binary(df.human), binary(df.judge)
    b_ = int(((jb == "PRESERVED") & (hb != "PRESERVED")).sum())
    c_ = int(((jb != "PRESERVED") & (hb == "PRESERVED")).sum())
    mcn = stats.binomtest(min(b_, c_), b_ + c_, 0.5).pvalue if b_ + c_ else float("nan")
    say(f"Direction of the Preserved disagreements: judge Preserved / human not {b_}, "
        f"human Preserved / judge not {c_} (exact McNemar p = {mcn:.3g}).")
    say()

    say("Disagreements by tool (judge -> human):")
    for t in TOOLS:
        sub = df[(df.method == t) & ~df.agree]
        pairs = sub.groupby(["judge", "human"]).size().sort_values(ascending=False)
        say(f"  {t:<18} {len(sub):>3} of {(df.method == t).sum()}: "
            + ", ".join(f"{j[0]}->{h[0]} {n}" for (j, h), n in pairs.items()))
    say("  (P = Preserved, O = Omitted, I = Incorrect)")
    say()

    # --- summary-level scores
    summ = (df.assign(hp=df.human == "PRESERVED", jp=df.judge == "PRESERVED",
                      hi=df.human == "INCORRECT", ji=df.judge == "INCORRECT")
              .groupby(["file_id", "project", "design_pattern", "file", "letter", "method"], as_index=False)
              .agg(n_facts=("fact_id", "size"), human_preserved_rate=("hp", "mean"),
                   judge_preserved_rate=("jp", "mean"), human_incorrect=("hi", "sum"),
                   judge_incorrect=("ji", "sum")))
    summ["difference"] = summ.judge_preserved_rate - summ.human_preserved_rate
    summ.to_csv(f"{OUT_PREFIX}_summaries.csv", index=False, float_format="%.4f")
    x, y = summ.human_preserved_rate.to_numpy(), summ.judge_preserved_rate.to_numpy()
    rho, rho_p = stats.spearmanr(x, y)
    r_, r_p = stats.pearsonr(x, y)
    icc = icc_a1(x, y)
    nonzero = summ.difference[summ.difference != 0]
    w_p = stats.wilcoxon(nonzero).pvalue if len(nonzero) else float("nan")
    say("2. Preserved rate per summary (120 summaries = 30 files x 4 tools)")
    say("-" * 78)
    say(f"Spearman rho = {rho:.2f} (p = {rho_p:.2g}); Pearson r = {r_:.2f} (p = {r_p:.2g}); ICC(A,1) = {icc:.2f}")
    say(f"Mean difference judge - human = {summ.difference.mean():+.3f} "
        f"(median {summ.difference.median():+.3f}; Wilcoxon p = {w_p:.3g}); "
        f"identical rate in {int((summ.difference == 0).sum())} of {len(summ)} summaries; "
        f"|difference| >= 0.5 in {int((summ.difference.abs() >= 0.5).sum())}.")
    for t in TOOLS:
        s = summ[summ.method == t]
        rho_t = stats.spearmanr(s.human_preserved_rate, s.judge_preserved_rate)[0]
        say(f"  {t:<18} mean diff {s.difference.mean():+.3f}   Spearman {rho_t:5.2f}   "
            f"ICC(A,1) {icc_a1(s.human_preserved_rate.to_numpy(), s.judge_preserved_rate.to_numpy()):5.2f}")
    say()

    # --- tool-level
    tool_rows = []
    for t in TOOLS:
        sub, s = df[df.method == t], summ[summ.method == t]
        row = {"method": t, "labels": len(sub)}
        for rater in ("human", "judge"):
            for lab in LABELS:
                row[f"{rater}_{lab.lower()}_rate"] = float((sub[rater] == lab).mean())
            row[f"{rater}_macro_preserved_rate"] = float(s[f"{rater}_preserved_rate"].mean())
        tool_rows.append(row)
    tools = pd.DataFrame(tool_rows)
    tools["human_rank"] = tools.human_preserved_rate.rank(ascending=False, method="min").astype(int)
    tools["judge_rank"] = tools.judge_preserved_rate.rank(ascending=False, method="min").astype(int)
    full = pd.read_csv(EVAL_DIR / "fact_preservation_summary.csv")
    full = full[(full.scope == "all")].set_index("method")
    tools["judge_preserved_rate_all_150"] = tools.method.map(full.preserved_rate)
    tools.to_csv(f"{OUT_PREFIX}_tools.csv", index=False, float_format="%.4f")
    say("3. Tools: label rates on the 30 files (micro: share of all facts)")
    say("-" * 78)
    say(f"{'Tool':<19}{'Preserved':>19}{'Omitted':>17}{'Incorrect':>17}{'Rank':>9}{'Judge 150':>11}")
    say(f"{'':<19}{'human  judge':>19}{'human  judge':>17}{'human  judge':>17}{'H  J':>9}{'files':>11}")
    for _, r in tools.iterrows():
        say(f"{r.method:<19}{r.human_preserved_rate:>11.1%}{r.judge_preserved_rate:>8.1%}"
            f"{r.human_omitted_rate:>9.1%}{r.judge_omitted_rate:>8.1%}"
            f"{r.human_incorrect_rate:>9.1%}{r.judge_incorrect_rate:>8.1%}"
            f"{r.human_rank:>6}{r.judge_rank:>3}{r.judge_preserved_rate_all_150:>11.1%}")
    say("Macro (mean of per-summary rates): " + "; ".join(
        f"{r.method} human {r.human_macro_preserved_rate:.1%} / judge {r.judge_macro_preserved_rate:.1%}"
        for _, r in tools.iterrows()))
    say()

    # --- per-file tool order
    pf_rows = []
    for fid, g in summ.groupby("file_id"):
        g = g.set_index("method").loc[TOOLS]
        h, j = g.human_preserved_rate.to_numpy(), g.judge_preserved_rate.to_numpy()
        same, opp, tie = pair_order(h, j)
        tau = stats.kendalltau(h, j).statistic if len(set(h)) > 1 and len(set(j)) > 1 else np.nan
        pf_rows.append({"file_id": fid, "file": g.file.iloc[0], "tau_b": tau, "pairs_same": same,
                        "pairs_opposite": opp, "pairs_tied": tie,
                        "human_best": ", ".join(g.index[h == h.max()]), "judge_best": ", ".join(g.index[j == j.max()]),
                        **{f"human_{t}": v for t, v in zip(TOOLS, h)}, **{f"judge_{t}": v for t, v in zip(TOOLS, j)}})
    pf = pd.DataFrame(pf_rows)
    pf.to_csv(f"{OUT_PREFIX}_per_file.csv", index=False, float_format="%.4f")
    untied = pf.pairs_same.sum() + pf.pairs_opposite.sum()
    say("4. Order of the four tools within each file (by preserved rate)")
    say("-" * 78)
    say(f"Kendall tau-b: mean {pf.tau_b.mean():.2f}, median {pf.tau_b.median():.2f} over "
        f"{pf.tau_b.notna().sum()} files (files where either rater ties all four tools are skipped).")
    say(f"Tool pairs ordered the same way: {pf.pairs_same.sum()} of {untied} pairs neither rater tied "
        f"({pf.pairs_same.sum() / untied:.0%}); opposite {pf.pairs_opposite.sum()}; "
        f"tied by at least one rater {pf.pairs_tied.sum()} of {6 * len(pf)}.")
    say(f"Same best tool (any overlap of ties): "
        f"{sum(bool(set(a.split(', ')) & set(b.split(', '))) for a, b in zip(pf.human_best, pf.judge_best))} "
        f"of {len(pf)} files.")
    say()

    # --- annotator-only
    ao_rows = []
    say("5. Annotator-only information (the judge has no counterpart)")
    say("-" * 78)
    fl = df.drop_duplicates(["file_id", "fact_id"])
    cc = fl.code_check.value_counts()
    say(f"Code check of the {len(fl)} extracted facts: " + ", ".join(f"{k} {v}" for k, v in cc.items()))
    for k, v in cc.items():
        ao_rows.append({"section": "code_check", "method": "", "value": k, "count": int(v)})
    bad = fl[fl.code_check != "Correct"]
    for _, r in bad.iterrows():
        say(f"  {r.file_id} fact {r.fact_id} ({r.fact_type}): {r.fact_text}  [{r.code_check}: "
            f"{(r.code_check_note or '').strip()}]")
    sub = df[df.code_check != "Correct"]
    if len(sub):
        say(f"  On these facts: agreement {np.mean(sub.human == sub.judge):.0%} over {len(sub)} labels.")
    say()
    say("Unsupported claims about the pattern (annotator, optional row; Yes = summary claims something "
        "the code does not support):")
    for t in TOOLS:
        c = claims[claims.method == t].unsupported_claim.value_counts()
        ao_rows += [{"section": "unsupported_claim", "method": t, "value": k, "count": int(v)} for k, v in c.items()]
        say(f"  {t:<18} Yes {c.get('Yes', 0):>2}   No {c.get('No', 0):>2}   blank {30 - c.sum():>2}")
    say()
    say(f"Extra facts recorded by the annotator: {extra.extra_id.nunique()} in "
        f"{extra.file_id.nunique()} files ("
        + ", ".join(f"{k} {v}" for k, v in extra.drop_duplicates('extra_id').fact_type.value_counts().items()) + ")")
    say(f"{'Tool':<19}{'Preserved':>10}{'Omitted':>9}{'Incorrect':>11}")
    for t in TOOLS:
        c = extra[extra.method == t].label.value_counts()
        ao_rows += [{"section": "extra_facts", "method": t, "value": k, "count": int(v)} for k, v in c.items()]
        say(f"{t:<19}{c.get('Preserved', 0):>10}{c.get('Omitted', 0):>9}{c.get('Incorrect', 0):>11}")
    role = extra[extra.fact_type == "Pattern role"]
    say(f"Of these, pattern-role extra facts ({role.extra_id.nunique()}), labelled Incorrect: " + ", ".join(
        f"{t} {int((role[role.method == t].label == 'Incorrect').sum())}" for t in TOOLS))
    pd.DataFrame(ao_rows).to_csv(f"{OUT_PREFIX}_annotator_only.csv", index=False)

    say()
    for line in latex_table(metrics, tools):
        say(line)
    write_block(RESULTS_TXT, "\n".join([BLOCK_START, *out, BLOCK_END]) + "\n")


if __name__ == "__main__":
    main()
